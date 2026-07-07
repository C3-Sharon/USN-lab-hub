from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("docs/planning")


TITLE_COLOR = RGBColor(31, 77, 120)
HEADING_COLOR = RGBColor(46, 116, 181)
MUTED = RGBColor(89, 89, 89)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_width(table):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), "9360")


def style_document(doc, title_text):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, HEADING_COLOR, 16, 8),
        ("Heading 2", 13, HEADING_COLOR, 12, 6),
        ("Heading 3", 12, TITLE_COLOR, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run(title_text)
    run.bold = True
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(20)
    run.font.color.rgb = TITLE_COLOR

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    srun = subtitle.add_run("USN Lab Hub IoT | 实验室硬件项目管理与智能运维平台")
    srun.font.name = "Calibri"
    srun._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    srun.font.size = Pt(10.5)
    srun.font.color.rgb = MUTED

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer.add_run("USN Lab Hub IoT 项目规划")
        footer_run.font.name = "Calibri"
        footer_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        footer_run.font.size = Pt(9)
        footer_run.font.color.rgb = MUTED


def add_p(doc, text, style=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(item)
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(item)
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.color.rgb = TITLE_COLOR
    paragraph2 = cell.add_paragraph()
    paragraph2.paragraph_format.space_after = Pt(0)
    run2 = paragraph2.add_run(body)
    run2.font.name = "Calibri"
    run2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.add_paragraph()


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        set_cell_text(header.cells[idx], text, bold=True)
        set_cell_shading(header.cells[idx], "F2F4F7")
    for row_data in rows:
        row = table.add_row()
        for idx, text in enumerate(row_data):
            set_cell_text(row.cells[idx], text)
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table


PROJECT_POSITIONING = (
    "在现有 USN-lab-hub 成员管理系统基础上，建设一个面向实验室硬件项目的轻量级 IoT 管理平台。"
    "平台以项目和设备资产为核心，通过 MQTT 接入真实硬件数据，实现设备状态监控、数据展示、"
    "异常告警、操作日志、远程指令与规则建议，最终形成从硬件接入到业务管理再到智能运维的完整闭环。"
)

WORKLOAD_RULE = (
    "项目构建节奏按每周 4-5 天安排，单人单日工作时长不超过 8 小时。"
    "周一用于需求和任务对齐，周三用于技术联调，周五或本周最后一个工作日用于阶段验收。"
)

SIX_WEEK_ROWS = [
    ["第 1 周", "共识、协议、表结构", "项目定位、MVP 边界、MQTT Topic、JSON 协议、ER 草案、页面原型", "架构评审"],
    ["第 2 周", "项目与设备资产管理", "项目/设备 CRUD、设备详情、标签状态、权限说明", "前后端字段联调"],
    ["第 3 周", "MQTT 接入与真实数据展示", "Broker、订阅客户端、原始/结构化数据入库、Redis 最新值、实时看板", "真实硬件上报验收"],
    ["第 4 周", "告警、日志、控制指令", "告警规则、指令下发、ACK 回执、操作日志、控制台", "业务闭环验收"],
    ["第 5 周", "规则建议与演示打磨", "建议生成、健康评分、访客展示、演示脚本、兜底模拟器", "完整演示彩排"],
    ["第 6 周", "测试、文档、答辩准备", "接口文档、部署说明、PPT、录屏、冻结功能、风险预案", "最终验收"],
]


def build_overall():
    doc = Document()
    style_document(doc, "USN Lab Hub IoT 总体项目发展规划")
    add_callout(doc, "项目共识", PROJECT_POSITIONING)

    doc.add_heading("1. 项目背景", level=1)
    add_p(doc, "USN-lab-hub 已完成成员登录、角色权限、成员管理、考勤管理等基础能力。与此同时，实验室硬件组长期积累了 IoT 边缘设备、功耗检测、音频巡检、智能水表、开发板、串行采集监测仪、智慧药盒、智慧超市等硬件项目。")
    add_p(doc, "这些硬件成果当前更像分散的项目资产，缺少统一的平台记录其负责人、所属项目、运行状态、数据指标、告警、维护记录和可演示成果。因此，本项目不应只做“传感器展示页”，而应建设一套围绕实验室硬件项目长期沉淀的管理与运维平台。")

    doc.add_heading("2. 项目定位与目标", level=1)
    add_p(doc, "项目定位：实验室硬件项目管理与智能运维平台。")
    add_p(doc, "一句话目标：在一个半月内完成真实硬件 MQTT 联动和完整业务闭环，并为后续设备规模化接入、规则引擎、实时推送、时序数据和 Agent 运维助手预留演进路径。")
    add_bullets(doc, [
        "用户范围：老师、实验室管理员、项目负责人、普通成员、硬件组成员、访客。",
        "核心闭环：人 -> 项目 -> 设备 -> 数据 -> 告警 -> 建议 -> 控制 -> 日志。",
        "验收重点：真实硬件联动、完整业务闭环、文档清晰、演示稳定。",
        WORKLOAD_RULE,
    ])

    doc.add_heading("3. 平台核心功能组成", level=1)
    add_table(doc, ["模块", "核心功能", "MVP 要求"], [
        ["用户与权限", "复用现有登录、角色、成员基础，扩展项目负责人、硬件负责人、访客视角", "基于角色控制菜单和操作权限"],
        ["项目管理", "管理实验室硬件项目、项目负责人、成员、项目状态、公开展示信息", "项目列表、详情、设备绑定"],
        ["设备资产", "记录硬件设备、模块、开发板、学生作品和客户项目", "设备 CRUD、状态、标签、负责人、所属项目"],
        ["MQTT 数据接入", "接入 telemetry/status/command/command_ack 消息", "至少 1-2 个真实设备上报成功"],
        ["数据展示", "展示最新值、历史数据、趋势图和设备在线状态", "设备详情实时卡片和历史曲线"],
        ["告警与建议", "规则判断异常，生成建议和处理记录", "阈值告警、离线告警、建议列表"],
        ["控制台", "通过 MQTT 下发低风险指令并接收 ACK", "至少一个可演示控制动作"],
        ["日志审计", "记录用户操作、数据接入、告警、指令、ACK", "关键链路可追踪"],
        ["访客展示", "展示公开项目成果和部分实时数据", "轻量展示页"],
    ], widths=[1.35, 3.1, 2.05])

    doc.add_heading("4. 技术路线", level=1)
    add_bullets(doc, [
        "前端：Vue 3、Element Plus、Vue Router、Axios、ECharts。",
        "后端：Spring Boot、MyBatis Plus、MySQL、Redis、Flyway、Swagger/OpenAPI、JWT。",
        "IoT 接入：优先使用 EMQX 或 Mosquitto 作为 MQTT Broker，后端使用 MQTT 客户端订阅和发布。",
        "MVP 不引入微服务、Kubernetes、复杂时序数据库和完整 Agent，以降低风险。",
        "后续根据数据量和实时性引入 WebSocket/SSE、消息队列、时序数据库和 Agent 服务。",
    ])

    doc.add_heading("5. MQTT 与数据协议规划", level=1)
    add_table(doc, ["消息类型", "Topic 规范", "用途"], [
        ["数据上报", "usn/{projectCode}/{deviceCode}/telemetry", "设备上报指标数据，如电压、电流、功率、温湿度等"],
        ["状态上报", "usn/{projectCode}/{deviceCode}/status", "设备在线、离线、维护、异常等状态"],
        ["指令下发", "usn/{projectCode}/{deviceCode}/command", "平台向设备发送低风险控制或配置指令"],
        ["指令回执", "usn/{projectCode}/{deviceCode}/command_ack", "硬件返回 ACK、失败原因或执行结果"],
    ], widths=[1.2, 3.25, 2.05])
    add_p(doc, "上报 JSON 建议包含 deviceCode、timestamp、metrics、status。指令 JSON 建议包含 commandId、command、params。ACK JSON 建议包含 commandId、status、message。")

    doc.add_heading("6. 数据模型规划", level=1)
    add_bullets(doc, [
        "lab_project：实验室硬件项目，连接成员、设备和展示信息。",
        "iot_device：设备资产，记录设备编号、名称、类型、协议、负责人、所属项目、状态。",
        "iot_device_metric：设备指标定义，记录 metric_key、显示名、单位、阈值和展示顺序。",
        "iot_telemetry_raw：原始 MQTT 消息，保留排错证据。",
        "iot_metric_data：解析后的结构化时序数据。",
        "iot_alert_record：告警记录，保存异常类型、等级、处理状态。",
        "iot_recommendation：规则建议或后续 Agent 建议。",
        "iot_command_record：指令下发和 ACK 状态。",
        "iot_operation_log：业务操作审计日志。",
    ])

    doc.add_heading("7. 一个半月实施路线", level=1)
    add_p(doc, WORKLOAD_RULE)
    add_table(doc, ["周期", "主题", "主要交付物", "里程碑"], SIX_WEEK_ROWS, widths=[0.85, 1.45, 3.3, 0.9])

    doc.add_heading("8. 三方对接机制", level=1)
    add_table(doc, ["时间", "会议目标", "必须产物"], [
        ["每周一", "需求与任务对齐，控制范围和风险", "本周任务清单、接口变更点、风险点"],
        ["每周三", "前后端和硬件技术联调", "联调记录、问题清单、接口或协议修订"],
        ["每周五/最后工作日", "阶段验收和演示检查", "可演示功能、截图/录屏、下周调整计划"],
        ["第 1 周末", "架构评审", "MVP 边界、ER 草案、MQTT 规范、首批硬件清单"],
        ["第 3 周末", "真实硬件联动验收", "硬件 MQTT 上报、后端入库、前端展示"],
        ["第 4 周末", "闭环验收", "告警、建议、指令、ACK、日志完整跑通"],
    ], widths=[1.2, 2.4, 2.9])

    doc.add_heading("9. 一个半月之后的发展路线", level=1)
    add_numbered(doc, [
        "扩展硬件项目资产库：接入更多开发板、检测模块、客户项目和学生作品，形成实验室长期资产沉淀。",
        "增强实时能力：引入 WebSocket 或 SSE，将 MQTT 数据变化主动推送到前端看板。",
        "增强规则能力：完善规则引擎、设备健康评分、趋势分析和告警处理流程。",
        "增强数据底座：在数据频率提高后，引入消息队列削峰，并评估 InfluxDB、TimescaleDB 等时序数据库。",
        "增强展示能力：建设项目大屏、设备拓扑、公开成果门户和移动端适配。",
        "建设 Agent 运维助手：读取设备数据、历史趋势、告警和日志，生成可解释建议；高风险控制必须保留人工确认。",
        "沉淀为实验室基础平台：将平台作为硬件项目立项、接入、展示、维护、复盘的统一入口。",
    ])

    doc.add_heading("10. 风险与兜底", level=1)
    add_table(doc, ["风险", "影响", "兜底方案"], [
        ["真实硬件不稳定", "演示中断", "保留 MQTT 模拟 Publisher，使用同一 Topic 和 JSON 协议"],
        ["控制动作有风险", "设备损坏或安全问题", "MVP 只做低风险动作，或做真实 ACK + 模拟执行"],
        ["时间不足", "功能散乱", "优先保证设备资产、MQTT 上报、看板、告警、指令、日志主链路"],
        ["数据格式频繁变化", "前后端和硬件联调成本上升", "第 1 周冻结第一版协议，后续变更必须记录"],
        ["Agent 目标过大", "拖慢 MVP", "一个半月内只做规则建议，Agent 放入后续路线"],
    ], widths=[1.45, 1.7, 3.35])

    doc.add_heading("11. 推荐演示路径", level=1)
    add_numbered(doc, [
        "管理员登录平台，进入实验室硬件项目总览。",
        "打开功耗检测或串行采集类设备详情。",
        "真实硬件通过 MQTT 上报功率、电压、电流或采集值。",
        "前端实时展示最新数据和历史曲线。",
        "系统根据阈值触发告警并生成处理建议。",
        "管理员下发低风险控制指令，例如修改采样周期或模拟开关。",
        "硬件返回 command_ack，平台更新指令状态。",
        "操作日志记录完整链路，访客页面展示公开项目成果。",
    ])

    return doc


def build_pm():
    doc = Document()
    style_document(doc, "产品经理任务说明与项目推进文档")
    add_callout(doc, "角色目标", "负责把项目讲清楚、边界控住、节奏推进，并把老师、前端、后端、硬件组的期望收敛到同一条业务闭环上。")

    doc.add_heading("1. 项目主线", level=1)
    add_p(doc, PROJECT_POSITIONING)
    add_p(doc, "产品经理需要持续提醒团队：本项目不是单纯的学生管理系统，也不是重型工业 IoT 平台，而是实验室硬件项目管理与智能运维平台。")

    doc.add_heading("2. 用户角色与核心场景", level=1)
    add_table(doc, ["用户", "主要诉求", "MVP 场景"], [
        ["老师", "了解实验室硬件成果、项目状态和平台建设路线", "查看总览和项目成果"],
        ["管理员", "管理成员、项目、设备、告警和日志", "完整闭环演示操作者"],
        ["项目负责人", "管理自己项目下的设备和数据", "查看项目设备和异常"],
        ["普通成员", "查看参与项目和设备状态", "浏览和协作"],
        ["硬件组", "按协议接入设备、验证数据和指令", "MQTT 上报与 ACK"],
        ["访客", "查看公开项目成果", "公开展示页"],
    ], widths=[1.15, 2.6, 2.75])

    doc.add_heading("3. 页面与需求清单", level=1)
    add_bullets(doc, [
        "首页总览：项目数量、设备数量、在线设备、告警数量、最近数据。",
        "项目列表与项目详情：项目介绍、负责人、成员、绑定设备、公开展示状态。",
        "设备列表与设备详情：设备编号、类型、协议、负责人、状态、标签、最新数据、历史曲线。",
        "告警中心：告警等级、设备、触发原因、处理状态、处理人。",
        "智能建议：基于规则生成处理建议，后续预留 Agent 来源。",
        "控制台：低风险指令下发、指令状态和 ACK 展示。",
        "操作日志：记录新增、修改、上报、告警、建议、指令、确认等动作。",
        "访客展示页：公开项目成果和部分设备数据。",
    ])

    doc.add_heading("4. 6 周推进计划", level=1)
    add_p(doc, WORKLOAD_RULE)
    add_table(doc, ["周期", "产品经理重点", "对接产物"], [
        ["第 1 周", "明确项目定位、用户角色、MVP 边界和首批硬件", "总需求草案、页面清单、演示闭环图、硬件接入清单"],
        ["第 2 周", "细化项目/设备字段和权限规则", "字段字典、设备状态定义、页面原型确认"],
        ["第 3 周", "推动真实 MQTT 联调，记录硬件指标和异常场景", "硬件接入说明、联调记录、数据指标表"],
        ["第 4 周", "定义告警、建议、控制动作和操作日志需求", "告警规则清单、控制指令清单、日志字段"],
        ["第 5 周", "固定演示故事线和答辩材料结构", "演示脚本、建议文案库、老师审阅版说明"],
        ["第 6 周", "组织彩排、冻结需求、完善材料", "PPT、测试记录、分工说明、最终验收清单"],
    ], widths=[0.85, 2.65, 3.0])

    doc.add_heading("5. 固定对接节奏", level=1)
    add_bullets(doc, [
        "周一：主持需求对齐，确认本周必须完成和可以延后的内容。",
        "周三：参加技术联调，记录接口、协议、硬件问题和责任人。",
        "周五或最后工作日：组织阶段验收，要求每个负责人拿出可展示产物。",
        "第 1 周末：确认 MVP 边界、第一批硬件、MQTT 协议和页面原型。",
        "第 3 周末：确认真实硬件数据可进入平台并展示。",
        "第 4 周末：确认告警、建议、指令、ACK、日志闭环跑通。",
    ])

    doc.add_heading("6. 与其他角色的对接产物", level=1)
    add_table(doc, ["对象", "产品经理提供", "产品经理接收"], [
        ["前端负责人", "页面原型、字段说明、状态文案、空状态和异常状态", "页面截图、交互问题、演示体验反馈"],
        ["后端负责人", "接口需求、权限规则、告警规则、建议文案、日志需求", "接口清单、数据字典、协议变更、风险说明"],
        ["硬件组", "首批硬件范围、指标表模板、演示场景", "设备编号、上报指标、单位、控制动作、联调结果"],
        ["老师", "总体规划、阶段成果、问题和风险", "方向反馈、验收偏好、功能优先级"],
    ], widths=[1.1, 2.7, 2.7])

    doc.add_heading("7. 一个半月后的产品路线", level=1)
    add_numbered(doc, [
        "把更多硬件项目纳入统一资产库，形成实验室成果沉淀。",
        "将访客展示页扩展为实验室公开项目门户。",
        "把规则建议升级为设备健康评分、趋势分析和维护工单。",
        "引入可配置看板，让老师和项目负责人自定义关注指标。",
        "建设 Agent 运维助手，但保持关键控制动作人工确认。",
    ])

    return doc


def build_frontend():
    doc = Document()
    style_document(doc, "前端负责人任务说明与页面建设文档")
    add_callout(doc, "角色目标", "负责把平台做成可看、可用、可演示的前端体验，让真实硬件数据、告警建议和控制闭环被用户清楚感知。")

    doc.add_heading("1. 技术栈与设计原则", level=1)
    add_bullets(doc, [
        "沿用现有 Vue 3、Element Plus、Vue Router、Axios 技术栈。",
        "新增 ECharts 用于历史数据和趋势图。",
        "页面风格应偏后台管理和数据看板，信息密度适中，避免做成营销页。",
        "优先保证演示链路稳定：总览 -> 项目 -> 设备 -> 数据 -> 告警 -> 控制 -> 日志。",
    ])

    doc.add_heading("2. 菜单与页面规划", level=1)
    add_table(doc, ["菜单", "页面", "优先级"], [
        ["IoT 总览", "项目数量、设备在线率、告警数量、最近数据", "P0"],
        ["项目管理", "项目列表、项目详情、项目绑定设备", "P0"],
        ["设备管理", "设备列表、设备详情、新增编辑、标签状态", "P0"],
        ["数据看板", "最新数据卡片、历史曲线、原始数据摘要", "P0"],
        ["告警中心", "告警列表、处理状态、设备筛选", "P0"],
        ["控制台", "指令按钮、参数输入、指令记录、ACK 状态", "P0"],
        ["操作日志", "按用户、设备、动作、时间查询", "P1"],
        ["访客展示", "公开项目和公开设备数据展示", "P1"],
    ], widths=[1.2, 4.0, 1.3])

    doc.add_heading("3. 数据展示规范", level=1)
    add_bullets(doc, [
        "设备状态使用明确颜色：在线、离线、异常、维护中。",
        "最新数据用卡片展示，历史数据用折线图或表格展示。",
        "告警需要突出等级、触发原因、处理状态和关联设备。",
        "控制指令必须展示状态流转：PENDING、SENT、ACKED、FAILED、TIMEOUT。",
        "真实硬件不可用时，前端不改协议，仍展示 MQTT 模拟器产生的数据。",
    ])

    doc.add_heading("4. 6 周任务安排", level=1)
    add_p(doc, WORKLOAD_RULE)
    add_table(doc, ["周期", "前端重点", "交付物"], [
        ["第 1 周", "搭建 IoT 菜单、路由、页面骨架，确认原型", "路由结构、菜单项、静态页面草图"],
        ["第 2 周", "完成项目和设备资产管理页面", "项目列表、设备列表、设备详情、新增编辑"],
        ["第 3 周", "完成数据看板和历史曲线", "最新数据卡片、ECharts 曲线、设备状态展示"],
        ["第 4 周", "完成告警中心和控制台", "告警列表、控制按钮、指令状态、操作日志入口"],
        ["第 5 周", "完成智能建议、访客展示和首页总览", "建议区域、公开展示页、总览大屏式首页"],
        ["第 6 周", "UI 打磨、演示录屏、异常兜底", "最终截图、演示视频、页面问题清单清零"],
    ], widths=[0.85, 2.7, 2.95])

    doc.add_heading("5. 与产品和后端的对接", level=1)
    add_table(doc, ["对接对象", "前端需要获得", "前端需要交付"], [
        ["产品经理", "页面原型、字段含义、状态文案、演示路径", "页面截图、交互反馈、缺失需求清单"],
        ["后端负责人", "Swagger/API 示例、错误码、字段枚举、鉴权规则", "接口联调记录、字段映射问题、前端期望返回结构"],
        ["硬件组", "通过后端获得稳定演示数据，不直接依赖硬件细节", "展示真实数据接入效果和控制 ACK 效果"],
    ], widths=[1.2, 2.7, 2.6])

    doc.add_heading("6. 一个半月后的前端路线", level=1)
    add_numbered(doc, [
        "引入 WebSocket/SSE，让实时数据主动刷新，减少轮询。",
        "建设可配置看板，支持不同角色保存关注设备和指标。",
        "建设项目公开门户，用于实验室成果展示和访客参观。",
        "增加设备拓扑或项目关系图，展示项目、设备、成员之间的关系。",
        "适配移动端，让现场调试和巡检更方便。",
    ])

    return doc


def build_backend():
    doc = Document()
    style_document(doc, "后端负责人任务说明与 IoT 接入设计文档")
    add_callout(doc, "角色目标", "负责把硬件数据可靠接进平台，把数据存住、查出、判断异常、下发指令并记录日志，形成可追踪的工程闭环。")

    doc.add_heading("1. 后端边界与技术栈", level=1)
    add_bullets(doc, [
        "沿用 Spring Boot、MyBatis Plus、MySQL、Redis、Flyway、Swagger、JWT。",
        "新增 MQTT 客户端能力，连接 EMQX 或 Mosquitto。",
        "MVP 先实现单体应用内的 IoT 模块，不拆微服务。",
        "数据设计以 device + metric 为核心，而不是只建 sensor 表，以兼容开发板、检测仪、智能设备和业务项目。",
    ])

    doc.add_heading("2. 模块规划", level=1)
    add_table(doc, ["模块", "职责", "MVP 接口/能力"], [
        ["项目模块", "管理实验室硬件项目", "项目 CRUD、项目绑定设备"],
        ["设备模块", "管理设备资产和状态", "设备 CRUD、状态更新、标签、负责人"],
        ["指标模块", "定义设备可上报指标", "metric_key、单位、阈值、显示名"],
        ["MQTT 接入", "订阅 telemetry/status，发布 command", "Broker 连接、消息解析、异常处理"],
        ["数据模块", "保存原始数据和结构化数据", "raw 入库、metric 入库、最新值 Redis"],
        ["告警模块", "规则判断异常", "阈值告警、离线告警"],
        ["建议模块", "基于规则生成建议", "recommendation 查询与状态更新"],
        ["指令模块", "下发指令并追踪 ACK", "command_record 状态机"],
        ["日志模块", "记录关键业务操作", "operation_log 查询"],
    ], widths=[1.25, 2.4, 2.85])

    doc.add_heading("3. 推荐数据库表", level=1)
    add_bullets(doc, [
        "lab_project：项目编号、项目名称、负责人、成员、状态、公开展示信息。",
        "iot_device：设备编号、设备名称、设备类型、通信协议、所属项目、负责人、状态、位置、描述。",
        "iot_device_metric：设备指标 key、名称、单位、阈值、显示顺序、是否启用。",
        "iot_telemetry_raw：Topic、payload、device_code、received_at、parse_status、error_message。",
        "iot_metric_data：device_id、metric_key、metric_value、unit、reported_at、received_at。",
        "iot_alert_record：设备、指标、等级、触发值、阈值、消息、状态、处理人。",
        "iot_recommendation：建议内容、来源 rule/agent、关联告警、状态。",
        "iot_command_record：command_id、device_id、command、params、status、sent_at、acked_at、message。",
        "iot_operation_log：操作人、动作、对象类型、对象 ID、摘要、时间。",
    ])

    doc.add_heading("4. MQTT 协议", level=1)
    add_table(doc, ["类型", "Topic", "后端处理"], [
        ["telemetry", "usn/{projectCode}/{deviceCode}/telemetry", "解析指标、写 raw、写 metric、更新 Redis 和 last_seen"],
        ["status", "usn/{projectCode}/{deviceCode}/status", "更新设备状态和在线时间"],
        ["command", "usn/{projectCode}/{deviceCode}/command", "后端发布指令给硬件"],
        ["command_ack", "usn/{projectCode}/{deviceCode}/command_ack", "更新 command_record 状态和日志"],
    ], widths=[1.1, 3.25, 2.15])
    add_p(doc, "建议上报 JSON：deviceCode、timestamp、metrics、status。建议指令 JSON：commandId、command、params。建议 ACK JSON：commandId、status、message。")

    doc.add_heading("5. API 规划", level=1)
    add_bullets(doc, [
        "项目：GET/POST/PUT /api/iot/projects，GET /api/iot/projects/{id}。",
        "设备：GET/POST/PUT /api/iot/devices，GET /api/iot/devices/{id}。",
        "数据：GET /api/iot/devices/{id}/latest，GET /api/iot/devices/{id}/metrics/history。",
        "告警：GET /api/iot/alerts，POST /api/iot/alerts/{id}/handle。",
        "建议：GET /api/iot/recommendations，POST /api/iot/recommendations/{id}/confirm。",
        "指令：POST /api/iot/devices/{id}/commands，GET /api/iot/commands。",
        "日志：GET /api/iot/operation-logs。",
    ])

    doc.add_heading("6. 6 周任务安排", level=1)
    add_p(doc, WORKLOAD_RULE)
    add_table(doc, ["周期", "后端重点", "交付物"], [
        ["第 1 周", "表结构、协议、接口草案", "Flyway 草案、MQTT Topic、JSON 示例、Swagger 分组"],
        ["第 2 周", "项目/设备/指标基础接口", "CRUD 接口、分页查询、状态枚举、操作日志基础"],
        ["第 3 周", "MQTT 接入和数据存储", "Broker 连接、订阅、raw 入库、metric 入库、Redis 最新值"],
        ["第 4 周", "告警、指令、ACK、日志", "规则扫描、command 发布、ACK 订阅、状态机、日志链路"],
        ["第 5 周", "建议、统计、演示兜底", "recommendation、健康评分、MQTT 模拟 Publisher、演示数据"],
        ["第 6 周", "测试、部署和文档", "接口文档、协议文档、部署说明、测试记录"],
    ], widths=[0.85, 2.6, 3.05])

    doc.add_heading("7. 与产品、前端、硬件组对接", level=1)
    add_table(doc, ["对象", "后端需要获得", "后端需要交付"], [
        ["产品经理", "字段定义、权限规则、告警规则、建议文案、演示设备", "接口清单、数据字典、协议风险、实现进度"],
        ["前端负责人", "页面字段、查询条件、交互状态、错误展示需求", "Swagger、API 示例、返回结构、错误码"],
        ["硬件组", "设备编号、指标名、单位、上报频率、控制动作、ACK 格式", "Broker 地址、Topic、JSON 协议、联调记录"],
    ], widths=[1.1, 2.8, 2.6])

    doc.add_heading("8. 风险控制", level=1)
    add_bullets(doc, [
        "每条 MQTT 消息先写 raw，再解析 metric，避免解析失败后无法排查。",
        "设备离线用 last_seen 判断，不依赖硬件主动声明。",
        "指令必须有 commandId，并设置超时状态，避免发出后无结论。",
        "真实控制只做低风险动作；高风险动作只生成建议或模拟 ACK。",
        "保留 MQTT 模拟 Publisher，确保演示时硬件异常也能跑通主链路。",
    ])

    doc.add_heading("9. 一个半月后的后端路线", level=1)
    add_numbered(doc, [
        "增加 WebSocket/SSE 推送最新数据和告警，减少前端轮询。",
        "引入消息队列缓冲高频上报，提升削峰和异步处理能力。",
        "评估时序数据库，支撑长期、高频、多指标历史数据查询。",
        "扩展规则引擎，支持更复杂的多指标判断和设备健康评分。",
        "接入 Agent 服务，让 Agent 读取数据、告警、日志并生成建议；关键指令仍需人工确认。",
        "完善权限审计、数据归档、部署自动化和监控告警。",
    ])

    return doc


def save_all():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    docs = [
        ("USN_Lab_Hub_IoT_总体项目发展规划.docx", build_overall()),
        ("USN_Lab_Hub_IoT_产品经理任务文档.docx", build_pm()),
        ("USN_Lab_Hub_IoT_前端负责人任务文档.docx", build_frontend()),
        ("USN_Lab_Hub_IoT_后端负责人任务文档.docx", build_backend()),
    ]
    for filename, doc in docs:
        doc.save(OUT_DIR / filename)


if __name__ == "__main__":
    save_all()
