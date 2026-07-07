from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("docs/planning")

TITLE_COLOR = RGBColor(31, 77, 120)
HEADING_COLOR = RGBColor(46, 116, 181)
MUTED = RGBColor(89, 89, 89)


def east_asia(run):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    east_asia(r)
    r.font.size = Pt(10)
    r.bold = bold


def set_table_width(table):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), "9360")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
        set_cell_shading(table.rows[0].cells[i], "F2F4F7")
    for row_data in rows:
        row = table.add_row()
        for i, text in enumerate(row_data):
            set_cell_text(row.cells[i], text)
            row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()


def style_document(doc, title):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, HEADING_COLOR, 16, 8),
        ("Heading 2", 13, HEADING_COLOR, 12, 6),
        ("Heading 3", 12, TITLE_COLOR, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    east_asia(r)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = TITLE_COLOR

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(14)
    sr = sub.add_run("USN Lab Hub IoT | 启动会议议程与线上协作手册")
    east_asia(sr)
    sr.font.size = Pt(10.5)
    sr.font.color.rgb = MUTED

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("USN Lab Hub IoT 启动会议与协作指南")
    east_asia(fr)
    fr.font.size = Pt(9)
    fr.font.color.rgb = MUTED


def add_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    east_asia(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        east_asia(r)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        east_asia(r)


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_width(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F7F7")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(lines):
        if idx:
            p.add_run("\n")
        r = p.add_run(line)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(9.5)
    doc.add_paragraph()


def build_doc():
    doc = Document()
    style_document(doc, "USN Lab Hub IoT 项目启动会议议程与协作指南")

    doc.add_heading("1. 文档用途", level=1)
    add_p(doc, "本文用于第一次线上启动会议。目标是让产品、前端、后端三位负责人在同一次会议中完成项目共识、代码拉取、Git 分支建立、GitLab/GitHub 同步策略确认、第一周任务分配和后续交接规则确认。")
    add_p(doc, "会议结束后，每个人都应清楚自己接下来做什么、产物交给谁、在哪里更新进度、代码如何合并、线上协作如何不断线。")

    doc.add_heading("2. 会前准备清单", level=1)
    add_table(doc, ["对象", "必须准备", "备注"], [
        ["所有人", "Git、IDE、实验室 GitLab 账号、项目仓库访问权限", "会议中要现场 clone、建分支、推送测试提交"],
        ["产品经理", "总规划文档、产品经理任务文档、硬件资料截图或表格", "用于确认第一批硬件和演示故事线"],
        ["前端负责人", "Node.js、npm 或 pnpm、浏览器调试环境", "会议中至少确认 frontend 可安装依赖和启动"],
        ["后端负责人", "JDK 17、Maven、MySQL、Redis、接口测试工具", "会议中至少确认 backend 运行条件"],
        ["硬件组对接人", "首批设备名单、数据指标、上报方式、可控动作", "如果不能参会，需提前提供表格"],
    ], widths=[1.2, 3.2, 2.1])

    doc.add_heading("3. 会议议程", level=1)
    add_table(doc, ["环节", "时间", "主持/负责人", "输出结果"], [
        ["项目共识展示", "10 分钟", "产品经理", "确认平台定位和业务闭环"],
        ["总体规划文档展示", "10 分钟", "产品经理", "确认 6 周路线和长期路线"],
        ["三人任务文档展示", "15 分钟", "三位负责人", "确认每人职责和第一周产物"],
        ["现有代码结构介绍", "10 分钟", "后端负责人", "明确 backend/frontend 目录和运行方式"],
        ["现场拉取项目", "20 分钟", "所有人", "每个人完成 clone 和本地分支创建"],
        ["GitLab/GitHub 策略确认", "10 分钟", "后端负责人", "确认 GitLab 主仓库、GitHub 镜像仓库"],
        ["任务看板建立", "15 分钟", "产品经理", "建立第一周任务清单和状态字段"],
        ["接口与 MQTT 表格建立", "15 分钟", "后端负责人", "建立 API 表、MQTT 协议表、设备接入表"],
        ["第一周任务确认", "15 分钟", "所有人", "明确负责人、交付物、对接时间"],
    ], widths=[1.45, 0.9, 1.35, 2.8])

    doc.add_heading("4. 推荐仓库策略", level=1)
    add_p(doc, "建议以实验室 GitLab 作为主仓库，GitHub 作为镜像、备份和作品展示仓库。日常开发、合并、Issue、MR 都在 GitLab 完成；GitHub 只同步稳定结果，避免两边同时改导致历史混乱。")
    add_table(doc, ["仓库", "定位", "使用规则"], [
        ["实验室 GitLab", "主仓库", "所有人日常开发、分支、Merge Request、任务管理均以 GitLab 为准"],
        ["GitHub", "镜像仓库", "由指定负责人阶段性同步 main/dev，用于展示和备份，不在 GitHub 直接改代码"],
    ], widths=[1.35, 1.55, 3.6])

    doc.add_heading("5. 分支规则", level=1)
    add_bullets(doc, [
        "main：稳定演示版，只放可以给老师看的版本。",
        "dev：日常集成版，前端、后端、文档功能合并到这里联调。",
        "feature/iot-product：产品经理维护需求、原型、会议纪要和文档。",
        "feature/iot-frontend：前端负责人维护 IoT 页面。",
        "feature/iot-backend：后端负责人维护 IoT 后端、MQTT、数据库和接口。",
        "feature/iot-mqtt-demo：MQTT 模拟器和硬件联调脚本，可由后端或硬件组维护。",
    ])
    add_p(doc, "feature 分支可以个人自由提交；dev 只能通过 Merge Request 合并；main 只能从 dev 合并稳定版本。每周最后一个工作日从 dev 选一次可演示版本合并到 main。")

    doc.add_heading("6. 首次拉取代码步骤", level=1)
    add_p(doc, "以下命令中的仓库地址需要替换为实验室 GitLab 的真实地址。第一次会议中每个人都应现场完成。")
    add_code(doc, [
        "git config --global user.name \"你的姓名或英文名\"",
        "git config --global user.email \"你的邮箱\"",
        "git clone <实验室GitLab仓库地址>",
        "cd USN-lab-hub",
        "git status",
    ])
    add_p(doc, "根据角色创建自己的功能分支：")
    add_code(doc, [
        "git checkout -b feature/iot-backend",
        "# 或者：git checkout -b feature/iot-frontend",
        "# 或者：git checkout -b feature/iot-product",
    ])
    add_p(doc, "完成一次测试提交，确认自己拥有推送权限：")
    add_code(doc, [
        "git status",
        "git add .",
        "git commit -m \"docs: add initial collaboration note\"",
        "git push origin feature/iot-backend",
    ])
    add_p(doc, "推送成功后，到 GitLab 页面创建 Merge Request，目标分支选择 dev。第一次 MR 可以只用于验证流程，不急着合并。")

    doc.add_heading("7. GitHub 同步方式", level=1)
    add_p(doc, "建议由后端负责人或项目统筹者统一负责 GitHub 同步。其他同学只需要对 GitLab 负责，减少误操作。")
    add_code(doc, [
        "git remote -v",
        "git remote add github <GitHub仓库地址>",
        "git push github main",
        "git push github dev",
    ])
    add_p(doc, "每周阶段验收后，同步一次 GitHub。不要在 GitHub 页面直接修改代码；如果发生 GitHub 和 GitLab 不一致，以 GitLab 为准。")

    doc.add_heading("8. 本地运行检查", level=1)
    add_table(doc, ["模块", "检查步骤", "通过标准"], [
        ["后端", "进入 backend，确认 JDK 17、Maven、MySQL、Redis 配置；启动 Spring Boot", "服务能启动，Swagger 或接口能访问"],
        ["前端", "进入 frontend，安装依赖并启动 Vite", "登录页或现有页面能打开"],
        ["数据库", "确认 MySQL 库存在，Flyway 可执行", "现有表能创建或迁移成功"],
        ["Redis", "确认 Redis 服务可连接", "登录态或缓存功能不报错"],
        ["MQTT", "第 3 周前搭建 EMQX/Mosquitto", "能发布和订阅测试消息"],
    ], widths=[1.0, 3.5, 2.0])

    doc.add_heading("9. 线上协作规则", level=1)
    add_bullets(doc, [
        "每周 4-5 天用于项目构建，单人单日工作时长不超过 8 小时。",
        "周一晚上开任务规划会，确认本周任务、依赖、交付物和风险。",
        "周三晚上开技术联调会，确认接口、MQTT、硬件和前端展示问题。",
        "周五或本周最后一个工作日开阶段验收会，必须拿出可展示产物。",
        "平时可以异步沟通，但重要结论必须写入会议纪要或任务表。",
        "每天结束前发简短日报，避免线上协作断线。",
    ])

    doc.add_heading("10. 每日简短日报模板", level=1)
    add_code(doc, [
        "今天完成：",
        "1. ",
        "2. ",
        "",
        "遇到问题：",
        "1. ",
        "",
        "明天计划：",
        "1. ",
        "",
        "需要谁配合：",
        "1. ",
    ])

    doc.add_heading("11. 交接产物模板", level=1)
    add_p(doc, "交接不要只靠聊天记录。建议固定维护三张表：接口交接表、MQTT 协议表、设备接入表。")
    add_table(doc, ["接口", "方法", "路径", "负责人", "状态", "前端联调"], [
        ["设备分页", "GET", "/api/iot/devices", "后端", "待开发", "否"],
        ["最新数据", "GET", "/api/iot/devices/{id}/latest", "后端", "待开发", "否"],
    ], widths=[1.1, 0.8, 2.3, 0.9, 0.9, 0.8])
    add_table(doc, ["类型", "Topic", "方向", "JSON 示例", "硬件支持"], [
        ["telemetry", "usn/{projectCode}/{deviceCode}/telemetry", "硬件 -> 平台", "deviceCode/timestamp/metrics/status", "待确认"],
        ["command_ack", "usn/{projectCode}/{deviceCode}/command_ack", "硬件 -> 平台", "commandId/status/message", "待确认"],
    ], widths=[1.0, 2.5, 1.1, 1.3, 0.6])
    add_table(doc, ["设备", "编号", "指标", "单位", "上报频率", "控制动作", "负责人"], [
        ["功耗检测", "PM-001", "voltage/current/power", "V/A/W", "待确认", "修改采样周期", "待确认"],
        ["串行采集数据监测仪", "DAQ-001", "analog_value", "待确认", "待确认", "待确认", "待确认"],
    ], widths=[1.2, 0.9, 1.7, 0.7, 0.9, 1.2, 0.7])

    doc.add_heading("12. 第一周任务清单", level=1)
    add_table(doc, ["负责人", "第一周任务", "交付物"], [
        ["产品经理", "确认首批硬件、用户角色、页面清单、演示故事线", "首批硬件接入清单、页面原型、演示流程草案"],
        ["前端负责人", "跑通前端、建立 IoT 菜单和页面骨架、用 mock 数据搭页面", "菜单路由、页面骨架、静态截图"],
        ["后端负责人", "跑通后端、设计表结构、制定 MQTT 协议和 API 草案", "ER 草案、Flyway 草案、MQTT 协议表、接口清单"],
    ], widths=[1.2, 3.2, 2.1])

    doc.add_heading("13. 合并与验收规则", level=1)
    add_bullets(doc, [
        "所有功能先合并到 dev，确认能启动后再考虑进入 main。",
        "合并前至少一名其他成员查看 MR，重点看是否破坏现有功能和接口约定。",
        "每周最后一个工作日冻结一次可演示版本，合并到 main 并同步 GitHub。",
        "如果发生冲突，由修改同一模块的人共同在线解决，不要用覆盖文件的方式处理。",
        "如果真实硬件临时不可用，使用 MQTT 模拟器跑同一套 Topic 和 JSON 协议。",
    ])

    doc.add_heading("14. 最终交付目录建议", level=1)
    add_code(doc, [
        "USN_Lab_Hub_IoT_项目交付/",
        "  01_项目规划文档/",
        "  02_产品原型与需求/",
        "  03_接口与MQTT协议/",
        "  04_数据库设计/",
        "  05_前端页面截图/",
        "  06_硬件联调记录/",
        "  07_测试与部署说明/",
        "  08_答辩PPT与演示脚本/",
    ])

    doc.add_heading("15. 会议结束确认项", level=1)
    add_numbered(doc, [
        "三位负责人都能访问 GitLab 仓库。",
        "三位负责人都完成本地 clone 和个人分支创建。",
        "至少一人完成 GitHub remote 配置，并明确同步负责人。",
        "任务看板已经建立，并录入第一周任务。",
        "接口交接表、MQTT 协议表、设备接入表已经建立。",
        "周一、周三、周五的固定会议时间已经确定。",
        "第一批硬件和第一条演示闭环已经初步确认。",
    ])

    return doc


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_doc().save(OUT_DIR / "USN_Lab_Hub_IoT_项目启动会议议程与协作指南.docx")
